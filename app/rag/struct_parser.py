"""
统一结构化解析器（非 PDF 链路）：docx / md / txt / pptx / html → schema v3 chunk

为什么存在这个模块
------------------
PDF 只能靠**几何推断**结构（所以有 1600 行 ``pdf_layout.py`` + 900 行
``pdf_structure.py``）；而 docx / md / pptx / html 的结构是**显式标注**的——
docx 的 ``paragraph.style.name`` 直接给出 ``Heading 1``、``document.tables`` 直接给
行列对象，md 的 ``#`` 数量就是层级，html 的 ``<h1>``/``<table>`` 就是结构。
让这些格式走"补齐默认值"的旧链路，等于**结构最好的文档拿到了最弱的结果**
（``section_path=[]`` / ``chunk_type`` 恒为 ``body`` / ``table=None``，且 docx 的
``document.tables`` 根本没被读取——表格内容整块丢失）。

本模块不做任何版面推断，只做「把已有结构读出来」。

两段式架构
----------
1. **适配器**：各格式 → 统一的中间表示 ``StructBlock`` 序列（标题 / 正文 / 表格）
2. **组装器**：``StructBlock`` 序列 → schema v3 chunk，字段与
   ``pdf_parser._make_chunk`` **逐字段对齐**

为什么中间表示不用"统一转 Markdown"
------------------------------------
转 md 会把表格压成管道文本、丢掉行列结构；而各格式的表格本来就是结构化的
（``document.tables`` / ``GraphicFrame`` / ``<table>``），压成文本再解析属于
无谓的有损往返。中间表示直接保留 ``kind`` / ``level`` / ``rows``，一次成型。

与 PDF 链路的差异（有意为之）
-----------------------------
- ``page_index`` / ``page_label`` 恒为 ``None``：这些格式没有"页"的概念
  （pptx 的 slide 更接近"页"，但语义不同，不冒充页码）；
- ``needs_ocr`` 恒为 ``False``：不存在扫描页；
- ``table`` 元数据不含 ``bbox``（无坐标），保留其余字段；
- ``parser_version`` 用 ``STRUCT_PARSER_VERSION`` 与 PDF 区分。
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from app.rag.chunker import _normalize_layout, _sanitize_text, split_text
from app.rag.config import STRUCT_PARSER_VERSION

# 表格单块最多数据行（与 ``pdf_structure.T_MAX_SPLIT_ROWS`` 对齐）：
# 超长表按行组切块、**每块重复表头**——表头是表格检索的命门
T_MAX_SPLIT_ROWS = 12


@dataclass
class StructBlock:
    """统一中间表示：一个结构单元（标题 / 正文段落 / 表格）"""

    kind: str  # "heading" | "body" | "table"
    text: str = ""  # heading 的标题文本 / body 的段落文本
    level: int | None = None  # heading 层级（1~6）
    rows: list[list[str]] | None = None  # table 的行列（含表头行）
    header_rows: int = 1  # 表头行数
    caption: str = ""  # 表格题注（可选）
    units: list[str] = field(default_factory=list)  # 表格单位（可选）


# --------------------------------------------------------------------------
# 表格线性化 / 切分（与 PDF 链路同一套文本形态，便于两路结果相互比较）
# --------------------------------------------------------------------------


def _table_to_text(
    block: StructBlock, rows: list[list[str]] | None = None, *, with_caption: bool = True
) -> str:
    """
    表格线性化成可检索文本

    格式与 ``pdf_structure.TableBlock.to_text`` 一致：题注 / 单位 / 表头 / 数据行，
    每行以 `` | `` 分隔。两路形态一致，检索与评测才能互相参照。
    """
    parts: list[str] = []
    if with_caption and block.caption:
        parts.append(block.caption)
    if block.units:
        parts.append("单位: " + " ".join(block.units))
    all_rows = block.rows or []
    for row in all_rows[: block.header_rows]:
        parts.append(" | ".join(row))
    body = all_rows[block.header_rows :] if rows is None else rows
    for row in body:
        parts.append(" | ".join(row))
    return "\n".join(p for p in parts if p.strip())


def _split_table(block: StructBlock) -> list[str]:
    """超长表按行组切块，每块重复表头；短表整块返回"""
    rows = block.rows or []
    if not rows:
        return []
    body = rows[block.header_rows :]
    if len(body) <= T_MAX_SPLIT_ROWS:
        return [_table_to_text(block)]
    return [
        _table_to_text(block, rows=body[i : i + T_MAX_SPLIT_ROWS])
        for i in range(0, len(body), T_MAX_SPLIT_ROWS)
    ]


def _collapse_cell(text: str) -> str:
    """单元格文本归一：换行压成空格、折叠连续空白"""
    return re.sub(r"\s+", " ", (text or "").replace("\n", " ")).strip()


# --------------------------------------------------------------------------
# 适配器 1：DOCX
# --------------------------------------------------------------------------

# 标题样式名：中文 Word 是「标题 1」，英文是「Heading 1」
_DOCX_HEADING_RE = re.compile(r"^\s*(?:heading|标题)\s*([1-9])\s*$", re.IGNORECASE)


def _iter_docx_items(document):
    """
    按**文档顺序**交替产出段落与表格

    ``document.paragraphs`` 与 ``document.tables`` 是两个独立列表，分别遍历会丢失
    二者在正文中的相对顺序（表格会被整体挪到段落之后）。这里直接走 XML 层的
    ``body`` 子元素，按出现顺序分发——这是 python-docx 官方给出的 recipe。
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _docx_heading_level(paragraph) -> int | None:
    """
    取段落的标题层级；不是标题则返回 ``None``

    两个信号，按可靠性排序：
    1. 样式名（``Heading 1`` / ``标题 1``）——显式套用了标题样式；
    2. 段落的 ``outlineLvl``（0-based）——部分文档不套标题样式，直接设大纲级别。
    """
    name = (paragraph.style.name or "").strip()
    match = _DOCX_HEADING_RE.match(name)
    if match:
        return int(match.group(1))
    pPr = getattr(paragraph._p, "pPr", None)
    if pPr is not None:
        from docx.oxml.ns import qn

        outline = pPr.find(qn("w:outlineLvl"))
        if outline is not None:
            val = outline.get(qn("w:val"))
            if val is not None and val.isdigit():
                return min(9, int(val) + 1)
    return None


def _docx_row_texts(row) -> list[str]:
    """
    取一行各单元格文本

    同一行内**按 XML 元素身份去重**：水平合并单元格会让 ``row.cells`` 把同一个
    ``tc`` 元素重复返回（合并了几列就重复几次）。按 ``id(cell._tc)`` 去重可精确
    还原成"合并后的那一格"，且不会误合"相邻但确实同值的两个独立单元格"。
    """
    seen: set[int] = set()
    texts: list[str] = []
    for cell in row.cells:
        key = id(cell._tc)
        if key in seen:
            continue
        seen.add(key)
        texts.append(_collapse_cell(cell.text))
    return texts


def blocks_from_docx(path: Path) -> list[StructBlock]:
    """DOCX → StructBlock 序列（段落按序成块，标题带层级，表格保留行列）"""
    import docx
    from docx.table import Table

    document = docx.Document(str(path))
    blocks: list[StructBlock] = []
    for item in _iter_docx_items(document):
        if isinstance(item, Table):
            rows = [_docx_row_texts(row) for row in item.rows]
            rows = [r for r in rows if any(c.strip() for c in r)]
            if rows:
                blocks.append(StructBlock(kind="table", rows=rows, header_rows=1))
            continue
        text = _collapse_cell(item.text)
        if not text:
            continue
        level = _docx_heading_level(item)
        if level is not None:
            blocks.append(StructBlock(kind="heading", text=text, level=level))
        else:
            blocks.append(StructBlock(kind="body", text=text))
    return blocks


# --------------------------------------------------------------------------
# 适配器 2：Markdown
# --------------------------------------------------------------------------

_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
# 表格分隔行：| --- | :---: | ---: |
_MD_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(?:\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _md_split_row(line: str) -> list[str]:
    """拆一行管道表：去掉首尾竖线后按 | 切分并 strip 每格"""
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def _md_table_start(lines: list[str], index: int) -> bool:
    """当前行含 ``|`` 且**下一行是分隔行**时，视为表格开始（否则只是普通含竖线文本）"""
    if "|" not in lines[index]:
        return False
    return index + 1 < len(lines) and bool(_MD_TABLE_SEP_RE.match(lines[index + 1]))


def blocks_from_markdown(path: Path) -> list[StructBlock]:
    """Markdown → StructBlock 序列（``#`` 层级成标题，管道表成表格，其余为正文）"""
    lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
    blocks: list[StructBlock] = []
    pending: list[str] = []
    index = 0
    total = len(lines)

    def flush() -> None:
        text = "\n".join(pending).strip()
        if text:
            blocks.append(StructBlock(kind="body", text=text))
        pending.clear()

    while index < total:
        line = lines[index]
        heading = _MD_HEADING_RE.match(line)
        if heading:
            flush()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if title:
                blocks.append(StructBlock(kind="heading", text=title, level=level))
            index += 1
            continue
        if _md_table_start(lines, index):
            flush()
            header = _md_split_row(line)
            index += 2  # 跳过表头行与分隔行
            body: list[list[str]] = []
            while index < total and lines[index].strip() and "|" in lines[index]:
                body.append(_md_split_row(lines[index]))
                index += 1
            rows = [header] + body
            blocks.append(StructBlock(kind="table", rows=rows, header_rows=1))
            continue
        pending.append(line)
        index += 1

    flush()
    return blocks


# --------------------------------------------------------------------------
# 适配器 3：纯文本
# --------------------------------------------------------------------------


def blocks_from_text(path: Path) -> list[StructBlock]:
    """
    纯文本 → StructBlock 序列

    按**空行分段**（原实现是"整篇一个块"，段落边界全丢，切分只能靠句号下钻）。
    """
    raw = path.read_text(encoding="utf-8", errors="ignore")
    blocks: list[StructBlock] = []
    for paragraph in re.split(r"\n\s*\n", raw):
        text = paragraph.strip()
        if text:
            blocks.append(StructBlock(kind="body", text=text))
    return blocks


# --------------------------------------------------------------------------
# 适配器 4：PPTX
# --------------------------------------------------------------------------


def _pptx_table_rows(table) -> list[list[str]]:
    """PPTX 表格 → 行列文本"""
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [_collapse_cell(cell.text) for cell in row.cells]
        if any(c for c in cells):
            rows.append(cells)
    return rows


def blocks_from_pptx(path: Path) -> list[StructBlock]:
    """
    PPTX → StructBlock 序列

    **每张 slide 产出一个一级标题**（取幻灯片标题；无标题时用「第 N 页」），使
    ``section_path`` 能定位到具体是哪一页——演示文稿是强分页的，这是它最自然的
    结构单元。slide 内的表格独立成块，其余文本框按 shape 顺序合成正文。

    跳过标题时比对 ``shape._element`` 而非 ``shape is title_shape``：python-docx /
    python-pptx 的包装对象**不是单例**，同一次迭代里 ``slide.shapes.title`` 与
    ``slide.shapes`` 取出的可能是两个不同的 Python 对象（指向同一 XML 元素），
    用 ``is`` 判断会漏掉标题、导致标题文本重复进正文。
    """
    from pptx import Presentation

    presentation = Presentation(str(path))
    blocks: list[StructBlock] = []
    for page, slide in enumerate(presentation.slides, start=1):
        try:
            title_shape = slide.shapes.title
        except Exception:  # 版式异常（无标题占位符）时降级为「第 N 页」
            title_shape = None
        title_element = title_shape._element if title_shape is not None else None
        title_text = _collapse_cell(title_shape.text) if title_shape is not None else ""

        blocks.append(
            StructBlock(kind="heading", text=title_text or f"第 {page} 页", level=1)
        )

        body_parts: list[str] = []

        def flush_body() -> None:
            if body_parts:
                blocks.append(StructBlock(kind="body", text="\n".join(body_parts)))
                body_parts.clear()

        for shape in slide.shapes:
            if title_element is not None and shape._element is title_element:
                continue
            if getattr(shape, "has_table", False):
                rows = _pptx_table_rows(shape.table)
                if rows:
                    flush_body()  # 表格与正文按 shape 顺序保序
                    blocks.append(StructBlock(kind="table", rows=rows, header_rows=1))
                continue
            text = _collapse_cell(getattr(shape, "text", "") or "")
            if text:
                body_parts.append(text)
        flush_body()
    return blocks


# --------------------------------------------------------------------------
# 适配器 5：HTML
# --------------------------------------------------------------------------


class _HtmlStructParser(HTMLParser):
    """
    把 HTML 解析成 ``StructBlock`` 序列

    只关心三类信号：``h1``~``h6``（标题 + 层级）、``table``（行列）、其余可见文本
    （正文）。``script`` / ``style`` / ``noscript`` 整体跳过。

    正文文本用 ``"".join`` 拼接而非加空格：行内标签（``<b>``、``<a>``）两侧的
    原始文本片段自带尾随空格（``"Hello "`` + ``"world"``），补空格反而会造出
    ``"Hello  world"`` 这类重复空白，交给 ``_normalize_layout`` 统一收尾即可。
    """

    _HEADING_LEVELS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}
    _SKIP_TAGS = {"script", "style", "noscript"}
    _BLOCK_TAGS = {"p", "div", "li", "section", "article", "br", "blockquote", "pre"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[StructBlock] = []
        self._skip_depth = 0
        self._heading_level: int | None = None
        self._buffer: list[str] = []
        self._table_rows: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    # ---- 收口 ----

    def _flush_body(self) -> None:
        text = _normalize_layout("".join(self._buffer))
        self._buffer = []
        if text:
            self.blocks.append(StructBlock(kind="body", text=text))

    def _flush_heading(self) -> None:
        text = _normalize_layout("".join(self._buffer))
        self._buffer = []
        if text and self._heading_level:
            self.blocks.append(
                StructBlock(kind="heading", text=text, level=self._heading_level)
            )
        self._heading_level = None

    def _flush_cell(self) -> None:
        if self._cell is not None and self._row is not None:
            self._row.append(_normalize_layout("".join(self._cell)))
        self._cell = None

    def _flush_row(self) -> None:
        self._flush_cell()
        if self._row is not None and any(c for c in self._row):
            if self._table_rows is not None:
                self._table_rows.append(self._row)
        self._row = None

    def _flush_table(self) -> None:
        self._flush_row()
        rows = self._table_rows or []
        self._table_rows = None
        if rows:
            self.blocks.append(StructBlock(kind="table", rows=rows, header_rows=1))

    # ---- HTMLParser 回调 ----

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag in self._HEADING_LEVELS:
            self._flush_body()
            self._flush_table()
            self._heading_level = self._HEADING_LEVELS[tag]
            self._buffer = []
        elif tag == "table":
            self._flush_body()
            self._table_rows = []
        elif tag == "tr":
            self._flush_row()
            self._row = []
        elif tag in ("td", "th"):
            self._flush_cell()
            self._cell = []
        elif tag in self._BLOCK_TAGS:
            self._flush_body()

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP_TAGS:
            if self._skip_depth:
                self._skip_depth -= 1
            return
        if self._skip_depth:
            return
        if tag in self._HEADING_LEVELS:
            self._flush_heading()
        elif tag in ("td", "th"):
            self._flush_cell()
        elif tag == "tr":
            self._flush_row()
        elif tag == "table":
            self._flush_table()
        elif tag in self._BLOCK_TAGS:
            self._flush_body()

    def handle_data(self, data: str) -> None:
        if self._skip_depth or not data:
            return
        if self._cell is not None:
            self._cell.append(data)
        else:
            self._buffer.append(data)

    def close(self) -> None:  # noqa: D102
        super().close()
        self._flush_heading()
        self._flush_body()
        self._flush_table()


def blocks_from_html(path: Path) -> list[StructBlock]:
    """HTML → StructBlock 序列"""
    parser = _HtmlStructParser()
    parser.feed(path.read_text(encoding="utf-8", errors="ignore"))
    parser.close()
    return parser.blocks


# --------------------------------------------------------------------------
# 分发
# --------------------------------------------------------------------------

_ADAPTERS = {
    ".docx": blocks_from_docx,
    ".md": blocks_from_markdown,
    ".markdown": blocks_from_markdown,
    ".txt": blocks_from_text,
    ".pptx": blocks_from_pptx,
    ".html": blocks_from_html,
    ".htm": blocks_from_html,
}

# 本模块能处理的全部后缀（与 chunker.SUPPORTED_SUFFIXES 取并集后即为总入口）
STRUCT_SUFFIXES = frozenset(_ADAPTERS)


def extract_blocks(file_path: Path) -> list[StructBlock]:
    """
    按后缀选择适配器，产出统一中间表示

    未知后缀返回空列表（与 ``chunker.extract_document`` 的约定一致）；适配器内部
    抛出的异常由调用方（``ingest``）统一降级处理。
    """
    adapter = _ADAPTERS.get(file_path.suffix.lower())
    if adapter is None:
        return []
    return adapter(Path(file_path))


def blocks_text(blocks: list[StructBlock]) -> str:
    """把中间表示拼成纯文本，供**文档级语言判定**使用（表格内容同样计入）"""
    parts: list[str] = []
    for block in blocks:
        if block.kind == "table":
            parts.append("\n".join(" ".join(row) for row in (block.rows or [])))
        elif block.text:
            parts.append(block.text)
    return "\n".join(p for p in parts if p)


# --------------------------------------------------------------------------
# 组装器：StructBlock 序列 → schema v3 chunk
# --------------------------------------------------------------------------


def _make_chunk(
    *,
    doc: str,
    text: str,
    section_path: list[str],
    section_index: int,
    seq: int,
    chunk_type: str,
    doc_lang: str,
    doc_cjk_ratio: float,
    table_meta: dict | None,
) -> dict:
    """
    组装单个 chunk（schema v3）

    字段与 ``pdf_parser._make_chunk`` 逐字段对齐，差异仅三处（见模块 docstring）：
    ``page_index`` / ``page_label`` 恒为 ``None``、``needs_ocr`` 恒为 ``False``、
    ``parser_version`` 用 ``STRUCT_PARSER_VERSION``。
    """
    text = _sanitize_text(text)
    heading = _sanitize_text(section_path[-1]) if section_path else ""
    section_path = [_sanitize_text(part) for part in section_path]
    if table_meta:
        table_meta = {
            k: (_sanitize_text(v) if isinstance(v, str) else v) for k, v in table_meta.items()
        }
    return {
        "id": f"{doc}::{section_index}::{heading or 'body'}::{seq}",
        "doc": doc,
        "heading": heading,
        "text": text,
        "lang": doc_lang,
        "cjk_ratio": doc_cjk_ratio,
        "zh_text": text if doc_lang == "zh" else "",
        "page_index": None,
        "page_label": None,
        "section_path": list(section_path),
        "chunk_type": chunk_type,
        "table": table_meta,
        "content_hash": hashlib.sha256(text.encode("utf-8")).hexdigest()[:16],
        "parser_version": STRUCT_PARSER_VERSION,
        "needs_ocr": False,
    }


def _renumber(chunks: list[dict]) -> None:
    """按（节序号, 出现顺序）重编块内序号，保证 id 唯一且稳定"""
    counters: dict[str, int] = {}
    for chunk in chunks:
        parts = chunk["id"].split("::")
        section = parts[1]
        counters[section] = counters.get(section, 0) + 1
        parts[-1] = str(counters[section])
        chunk["id"] = "::".join(parts)


def assemble(
    blocks: list[StructBlock], rel_name: str, doc_lang: str, doc_cjk_ratio: float
) -> list[dict]:
    """
    ``StructBlock`` 序列 → schema v3 chunk 列表

    装配规则与 ``pdf_parser.parse_document`` 一致：
    1. 标题入**标题栈**（level 小的压栈，同级 / 更高级弹栈），标题文本进
       ``section_path``、**不进正文**（``heading`` 是独立字段）；
    2. 表格独立成块（``chunk_type="table"``），超长表按行组切、每块重复表头；
    3. 连续正文累积到同一缓冲、按 ``split_text`` 的句子边界切块；
    4. 遇标题 / 表格 / 收尾时收口正文缓冲。
    """
    chunks: list[dict] = []
    stack: list[tuple[int, str]] = []
    section_path: list[str] = []
    section_index = 0
    seq = 0
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer, seq
        if not buffer:
            return
        # buffer 里各元素之间是**语义段落边界**（每个 body block 入队时已逐段做过
        # _normalize_layout）。这里**不再对整体调一次 _normalize_layout**——它会把
        # 中文之间的换行直接删掉（那是为 PDF 的视觉行设计的），用它收口会把相邻
        # 段落粘成一句。
        text = "\n".join(buffer).strip()
        buffer = []
        if not text:
            return
        for piece in split_text(text):
            seq += 1
            chunks.append(
                _make_chunk(
                    doc=rel_name,
                    text=piece,
                    section_path=section_path,
                    section_index=section_index,
                    seq=seq,
                    chunk_type="body",
                    doc_lang=doc_lang,
                    doc_cjk_ratio=doc_cjk_ratio,
                    table_meta=None,
                )
            )

    for block in blocks:
        if block.kind == "heading":
            flush()
            level = block.level or 1
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, block.text))
            section_path = [text for _, text in stack]
            section_index += 1
            seq = 0
        elif block.kind == "table":
            flush()
            pieces = _split_table(block)
            rows = block.rows or []
            for index, piece in enumerate(pieces, start=1):
                text = piece if len(pieces) == 1 else f"{piece}\n(第 {index}/{len(pieces)} 段)"
                seq += 1
                chunks.append(
                    _make_chunk(
                        doc=rel_name,
                        text=text,
                        section_path=section_path,
                        section_index=section_index,
                        seq=seq,
                        chunk_type="table",
                        doc_lang=doc_lang,
                        doc_cjk_ratio=doc_cjk_ratio,
                        table_meta={
                            "caption": block.caption,
                            "units": block.units,
                            "n_rows": len(rows),
                            "n_cols": max((len(row) for row in rows), default=0),
                            "header_rows": block.header_rows,
                            "bbox": None,
                            "part": index,
                            "parts": len(pieces),
                        },
                    )
                )
        elif block.text.strip():
            # 段内归一：txt / md / pptx 的段落里可能含手工换行或视觉换行，先合并；
            # 段与段之间的边界由上一步的 "\n".join 保留
            buffer.append(_normalize_layout(block.text))

    flush()
    _renumber(chunks)
    return chunks


def parse_document(
    file_path: Path | str,
    rel_name: str,
    *,
    doc_lang: str = "",
    doc_cjk_ratio: float = 0.0,
) -> list[dict]:
    """
    解析单份非 PDF 文档，产出 schema v3 的 chunk 列表

    入参与 ``pdf_parser.parse_document`` 对齐（除 PDF 专属的 ``use_text_tables``）。
    """
    blocks = extract_blocks(Path(file_path))
    return assemble(blocks, rel_name, doc_lang, doc_cjk_ratio)
